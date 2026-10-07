import urllib.request
import openpyxl
from openpyxl import load_workbook
import io
import json

url = 'https://docs.google.com/spreadsheets/d/e/2PACX-1vRXJKvpqfkPBpL9_TCR55PGZcybhS16TYuqQ6Phrnd3Bkapj2LrWcdV6OQdSrEOrA/pub?output=xlsx'
response = urllib.request.urlopen(url)
file_content = response.read()

wb = load_workbook(filename=io.BytesIO(file_content))
ws = wb.active

faculty = []
for row in ws.iter_rows(min_row=2):
    name_cell = row[1]
    if not name_cell.value:
        continue
    name = str(name_cell.value).strip()
    
    designation = str(row[2].value).strip() if row[2].value else ''
    joiningDate = str(row[3].value).strip() if row[3].value else ''
    qualification = str(row[4].value).strip() if row[4].value else ''
    association = str(row[5].value).strip() if row[5].value else ''
    
    profileUrl = ''
    link_cell = row[6]
    if link_cell.hyperlink:
        profileUrl = link_cell.hyperlink.target
    
    faculty.append({
        "name": name,
        "designation": designation,
        "qualification": qualification,
        "joiningDate": joiningDate,
        "association": association,
        "profileUrl": profileUrl
    })

with open('has_faculty_with_links.json', 'w') as f:
    json.dump(faculty, f, indent=4)
print("Saved links")
